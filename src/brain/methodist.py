"""Methodist Agent: Specialized educational and methodological AI agent for SUAI Physics Dept.

Implements 'Frozen Context' architecture with ephemeral Table of Contents (TOC) page slicing.
"""

import hashlib
import json
import logging
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pypdf
from google import genai
from google.genai import types

from src import config
from src.storage.db import Database
from src.persona.prompts import render_prompt

logger = logging.getLogger(__name__)


class MethodistService:
    """Core autonomous agent for academic regulations, GOSTs, textbooks, and curriculum.
    
    Implements 'Frozen Context' architecture:
    - Base context is clean, immutable and contains only system rules and brief catalog metadata.
    - Full-text textbook slices are loaded ephemerally on-the-fly and immediately discarded after response generation.
    - Separate training / re-indexing mode for populating TOC and FTS tables without polluting active context.
    """

    def __init__(
        self,
        db: Database,
        raw_dir: str | Path = "/opt/assistant-core/data/methodist/raw",
        kb_dir: str | Path = "/opt/assistant-core/data/methodist/knowledge_base",
        api_key: str | None = config.GEMINI_API_KEY,
        model: str = config.GEMINI_MODEL,
    ) -> None:
        self.db = db
        self.raw_dir = Path(raw_dir)
        self.kb_dir = Path(kb_dir)
        self.api_key = api_key or ""
        self.model = model
        self._client: genai.Client | None = None

    def get_client(self) -> genai.Client:
        if self._client is None:
            if not self.api_key:
                raise ValueError("GEMINI_API_KEY is not configured.")
            self._client = genai.Client(api_key=self.api_key)
        return self._client

    # --- 1. TOC Extraction & Document Parsing (Training Mode) ---

    def _extract_toc_from_pdf(self, pdf_path: Path) -> tuple[int, list[dict[str, Any]]]:
        """Extract Table of Contents (TOC) and page boundaries from a PDF document."""
        reader = pypdf.PdfReader(str(pdf_path))
        total_pages = len(reader.pages)
        entries: list[dict[str, Any]] = []

        # A. Try PDF outline (bookmarks)
        has_usable_outline = False
        if reader.outline:
            raw_items: list[dict[str, Any]] = []

            def walk_outline(items: list[Any], parent: str | None = None) -> None:
                for item in items:
                    if isinstance(item, list):
                        walk_outline(item, parent=raw_items[-1]["title"] if raw_items else None)
                    elif isinstance(item, pypdf.generic.Destination):
                        try:
                            p_num = reader.get_destination_page_number(item)
                        except Exception:
                            p_num = None
                        title = (item.title or "").strip()
                        # Filter out synthetic MS Word bookmarks
                        if title and not re.match(r"^(?:_GoBack|MTBlankEqn|bookmark\d+)$", title, re.IGNORECASE):
                            raw_items.append({
                                "title": title,
                                "page": p_num + 1 if p_num is not None else 1,
                                "parent": parent,
                            })

            walk_outline(reader.outline)
            # If we found at least 2 non-synthetic items
            if len(raw_items) >= 2:
                has_usable_outline = True
                for i, item in enumerate(raw_items):
                    p_start = item["page"]
                    if i + 1 < len(raw_items):
                        p_next = raw_items[i + 1]["page"]
                        p_end = max(p_start, p_next - 1) if p_next >= p_start else p_start
                    else:
                        p_end = total_pages

                    m = re.match(r"^(?:§|Глава|Раздел|Лабораторная работа №?\s*|)(\d+(?:\.\d+)*)", item["title"])
                    ch_num = m.group(1) if m else None

                    entries.append({
                        "chapter_title": item["title"],
                        "chapter_number": ch_num,
                        "page_start": p_start,
                        "page_end": p_end,
                        "parent_title": item["parent"],
                        "summary": None,
                    })

        # B. Fallback to printed TOC if bookmarks are missing or synthetic
        if not has_usable_outline or len(entries) < 3:
            printed_entries: list[dict[str, Any]] = []
            pages_to_check = list(range(min(8, total_pages))) + list(range(max(0, total_pages - 8), total_pages))
            for p_idx in pages_to_check:
                txt = reader.pages[p_idx].extract_text() or ""
                if "содержание" in txt.lower() or "оглавление" in txt.lower():
                    # Combine multi-line titles with dots and page numbers
                    current_title_parts: list[str] = []
                    for line in txt.split("\n"):
                        line_s = line.strip()
                        if not line_s or line_s.isdigit():
                            continue
                        m = re.search(r"^(.*?)[.\s…_-]{3,}\s*(\d+)\s*$", line_s)
                        if m:
                            part1 = m.group(1).strip()
                            page_num = int(m.group(2))
                            full_title = " ".join(current_title_parts + ([part1] if part1 else [])).strip()
                            if full_title and page_num <= total_pages:
                                m_num = re.match(r"^(?:§|Глава|Раздел|Лабораторная работа №?\s*|)(\d+(?:\.\d+)*)", full_title)
                                printed_entries.append({
                                    "chapter_title": full_title,
                                    "chapter_number": m_num.group(1) if m_num else None,
                                    "page_start": page_num,
                                    "page_end": page_num,
                                    "parent_title": None,
                                    "summary": None,
                                })
                            current_title_parts = []
                        else:
                            if not any(k in line_s.lower() for k in ["содержание", "оглавление"]):
                                current_title_parts.append(line_s)

                    if len(printed_entries) >= 2:
                        for i in range(len(printed_entries)):
                            if i + 1 < len(printed_entries):
                                printed_entries[i]["page_end"] = max(
                                    printed_entries[i]["page_start"],
                                    printed_entries[i + 1]["page_start"] - 1,
                                )
                            else:
                                printed_entries[i]["page_end"] = total_pages
                        entries = printed_entries
                        break

        # C. Default entry if no TOC was found
        if not entries:
            entries.append({
                "chapter_title": pdf_path.stem.replace("_", " "),
                "chapter_number": "1",
                "page_start": 1,
                "page_end": total_pages,
                "parent_title": None,
                "summary": None,
            })

        return total_pages, entries

    def _extract_docx_chunks(self, docx_path: Path) -> list[dict[str, str]]:
        """Parse DOCX document into structured chunks (paragraphs + tables)."""
        import docx

        doc = docx.Document(str(docx_path))
        chunks: list[dict[str, str]] = []

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            chunks.append({
                "section_title": "Текстовые реквизиты и формулировки",
                "chunk_text": "\n".join(paragraphs),
            })

        table_rows: list[str] = []
        for t_idx, table in enumerate(doc.tables):
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    table_rows.append(" | ".join(cells))

        if table_rows:
            chunks.append({
                "section_title": "Табличные формы и бланки измерений",
                "chunk_text": "\n".join(table_rows),
            })

        return chunks

    def _split_markdown_into_sections(self, content: str, default_title: str) -> list[dict[str, str]]:
        """Split markdown content into logical chunks based on headers."""
        lines = content.split("\n")
        chunks: list[dict[str, str]] = []
        current_title = default_title
        current_lines: list[str] = []

        for line in lines:
            if line.startswith("## ") or line.startswith("# "):
                if current_lines:
                    text_block = "\n".join(current_lines).strip()
                    if len(text_block) > 40:
                        chunks.append({"section_title": current_title, "chunk_text": text_block})
                    current_lines = []
                current_title = line.lstrip("#").strip()
            current_lines.append(line)

        if current_lines:
            text_block = "\n".join(current_lines).strip()
            if len(text_block) > 40:
                chunks.append({"section_title": current_title, "chunk_text": text_block})

        return chunks

    async def ingest_pdf_book(self, file_path: str | Path, category: str | None = None) -> dict[str, Any]:
        """Ingest a PDF textbook, lab manual, or problem book into TOC index without loading full text."""
        p = Path(file_path)
        if not p.is_file():
            raise FileNotFoundError(f"PDF file not found: {p}")

        content_hash = hashlib.sha256(p.read_bytes()).hexdigest()

        if not category:
            parent_name = p.parent.name
            category = parent_name if parent_name in ("textbooks", "problem_books", "lab_manuals") else "textbooks"

        title = p.stem.replace("_", " ")
        total_pages, toc_entries = self._extract_toc_from_pdf(p)

        summary = f"Учебное издание ГУАП ({category}). Всего страниц: {total_pages}. Разделов в оглавлении: {len(toc_entries)}."

        doc_id = await self.db.save_methodist_document(
            category=category,
            title=title,
            file_path=str(p.resolve()),
            content_hash=content_hash,
            summary=summary,
            key_terms=f"{category}, {title}",
            total_chunks=len(toc_entries),
        )

        for entry in toc_entries:
            entry["book_title"] = title
            entry["category"] = category

        saved_toc_count = await self.db.save_methodist_toc_entries(doc_id, toc_entries)

        # Also store brief catalog chunk into methodist_chunks for general search
        toc_overview = "\n".join([f"• [{e['page_start']}-{e['page_end']}] {e['chapter_title']}" for e in toc_entries[:25]])
        await self.db.replace_methodist_chunks(doc_id, [{
            "section_title": f"Оглавление: {title}",
            "chunk_text": f"Пособие: {title} ({total_pages} с.)\nОсновные разделы:\n{toc_overview}",
        }])

        logger.info("Ingested PDF book '%s' (ID %d) with %d TOC entries", title, doc_id, saved_toc_count)
        return {
            "doc_id": doc_id,
            "title": title,
            "category": category,
            "total_pages": total_pages,
            "toc_entries": saved_toc_count,
        }

    async def ingest_docx_template(self, file_path: str | Path, category: str = "templates") -> dict[str, Any]:
        """Ingest a DOCX title page or protocol blank into full-text chunks."""
        p = Path(file_path)
        if not p.is_file():
            raise FileNotFoundError(f"DOCX template not found: {p}")

        content_hash = hashlib.sha256(p.read_bytes()).hexdigest()
        title = p.stem.replace("_", " ")

        chunks = self._extract_docx_chunks(p)
        summary = f"Шаблон/бланк ГУАП: {title}. Содержит обязательные реквизиты, формулировки и формы таблиц."

        doc_id = await self.db.save_methodist_document(
            category=category,
            title=title,
            file_path=str(p.resolve()),
            content_hash=content_hash,
            summary=summary,
            key_terms=f"{category}, бланк, протокол, титульный лист, {title}",
            total_chunks=len(chunks),
        )

        inserted_chunks = await self.db.replace_methodist_chunks(doc_id, chunks)
        logger.info("Ingested DOCX template '%s' (ID %d) with %d chunks", title, doc_id, inserted_chunks)
        return {
            "doc_id": doc_id,
            "title": title,
            "category": category,
            "chunks_count": inserted_chunks,
        }

    async def ingest_markdown_file(self, file_path: str | Path, category: str | None = None) -> dict[str, Any]:
        """Ingest a Markdown regulation, GOST, or RPD file."""
        p = Path(file_path)
        if not p.is_file():
            raise FileNotFoundError(f"File not found: {p}")

        content = p.read_text(encoding="utf-8")
        content_hash = hashlib.sha256(content.encode("utf-8")).hexdigest()

        if not category:
            parent_name = p.parent.name
            category = parent_name if parent_name in ("gost", "rpd", "admin_lna", "dept_physics") else "general"

        title = p.stem.replace("_", " ")
        first_line = content.strip().split("\n")[0]
        if first_line.startswith("# "):
            title = first_line.lstrip("#").strip()

        chunks = self._split_markdown_into_sections(content, default_title=title)
        paras = [para.strip() for para in content.split("\n\n") if para.strip() and not para.strip().startswith("#")]
        summary_text = paras[0][:400] if paras else title

        doc_id = await self.db.save_methodist_document(
            category=category,
            title=title,
            file_path=str(p.resolve()),
            content_hash=content_hash,
            summary=summary_text,
            key_terms=category,
            total_chunks=len(chunks),
        )

        inserted_chunks = await self.db.replace_methodist_chunks(doc_id, chunks)
        logger.info("Ingested markdown '%s' (ID %d) with %d chunks", title, doc_id, inserted_chunks)
        return {
            "doc_id": doc_id,
            "title": title,
            "category": category,
            "chunks_count": inserted_chunks,
        }

    async def reindex_all_materials(self, force: bool = False) -> dict[str, Any]:
        """Scans all folders in data/methodist/raw/ and re-indexes Markdown, DOCX, and PDF books."""
        self.raw_dir.mkdir(parents=True, exist_ok=True)
        results: list[dict[str, Any]] = []

        # 1. Index Markdown files (GOST, RPD, LNA, Rules)
        for f in self.raw_dir.rglob("*.md"):
            try:
                res = await self.ingest_markdown_file(f)
                results.append(res)
            except Exception as e:
                logger.error("Error ingesting %s: %s", f.name, e)

        # 2. Index DOCX templates (Title sheets, Measurement protocol blanks)
        for f in self.raw_dir.rglob("*.docx"):
            try:
                res = await self.ingest_docx_template(f)
                results.append(res)
            except Exception as e:
                logger.error("Error ingesting template %s: %s", f.name, e)

        # 3. Index PDF textbooks, manuals, and problem books into TOC index
        for f in self.raw_dir.rglob("*.pdf"):
            try:
                res = await self.ingest_pdf_book(f)
                results.append(res)
            except Exception as e:
                logger.error("Error indexing PDF book %s: %s", f.name, e)

        return {
            "total_processed": len(results),
            "items": results,
        }

    # --- 2. Ephemeral Page Loading & Frozen Context Engine ---

    def get_ephemeral_pages(self, file_path: str, page_start: int, page_end: int, max_pages: int = 8) -> str:
        """Extract a targeted page slice from a PDF on disk.
        
        This slice is ephemeral: it is loaded into memory only for the duration
        of a single generation call and is never saved into persistent context.
        """
        p = Path(file_path)
        if not p.is_file():
            return f"[Файл {p.name} не найден на диске]"

        try:
            reader = pypdf.PdfReader(str(p))
            total_pages = len(reader.pages)
            actual_start = max(1, min(page_start, total_pages))
            actual_end = min(total_pages, max(actual_start, page_end), actual_start + max_pages - 1)

            page_texts = []
            for p_num in range(actual_start, actual_end + 1):
                txt = (reader.pages[p_num - 1].extract_text() or "").strip()
                if txt:
                    page_texts.append(f"--- [Стр. {p_num} из {total_pages}] ---\n{txt}")

            return "\n\n".join(page_texts) if page_texts else "[Текст на указанных страницах пуст или является изображением]"
        except Exception as e:
            logger.error("Failed to extract ephemeral pages from %s: %s", file_path, e)
            return f"[Ошибка чтения страниц: {e}]"

    # --- 3. Scope Verification & Query Answering ---

    async def check_scope(self, question: str) -> dict[str, Any]:
        """Check if question falls strictly inside Methodist domain using heuristics and LLM."""
        q_lower = question.lower()
        out_of_scope_keywords = [
            "тидекс", "табель", "смен", "часы работы", "выработка", "погод",
            "анекдот", "фильм", "игра", "криптовалют", "биткоин", "ресторан",
            "билет", "поезд", "рейс", "курс валют", "новости", "политик",
        ]
        if any(k in q_lower for k in out_of_scope_keywords):
            return {
                "is_in_scope": False,
                "category": "out_of_scope",
                "reason": "Вопрос относится к сторонней теме вне методической работы кафедры физики ГУАП",
            }

        academic_keywords = [
            "гост", "титульн", "отчет", "лабораторн", "физик", "фонт", "опд",
            "брс", "баллы", "погрешност", "протокол", "допуск", "защита",
            "практикум", "зачет", "экзамен", "кафедр", "попов", "гастелло",
            "формул", "график", "мнк", "шрифт", "поля", "студент", "семестр",
            "дедлайн", "пересдач", "рпд", "проектн", "лаба", "лабы", "методичк",
            "оформлен", "подпис", "правил", "коваленко", "задачник", "задач",
            "учебник", "комптон", "фотоэффект", "спектрометр", "ньютон",
            "интерференц", "дифракц", "маятник", "гистерезис", "холл",
            "эффективн", "контракт", "стимулирующ", "надбавк", "критери",
            "базов уровень", "повышенн уровень", "ппс", "новиков", "приказ 05-259",
            "печ.л", "п.л", "печатн лист", "соавтор", "варост",
        ]
        if any(k in q_lower for k in academic_keywords):
            return {"is_in_scope": True, "category": "academic", "reason": "Academic keyword matched"}

        try:
            client = self.get_client()
            prompt = render_prompt("methodist_scope_check", question=question)
            resp = await client.aio.models.generate_content(
                model=self.model,
                contents=prompt,
                config=types.GenerateContentConfig(
                    temperature=0.0,
                    response_mime_type="application/json",
                ),
            )
            raw = (resp.text or "{}").strip()
            return json.loads(raw)
        except Exception as e:
            logger.debug("Scope check error: %s", e)
            return {"is_in_scope": False, "category": "out_of_scope", "reason": "Unrecognized scope"}

    async def ask(
        self,
        question: str,
        caller: str = "user",
        caller_id: str | None = None,
    ) -> dict[str, Any]:
        """Ask a question to the Methodist Agent with strict scope guardrails and Ephemeral TOC Loading."""
        cleaned_q = question.strip()
        if not cleaned_q:
            return {
                "answer": "Пожалуйста, задайте вопрос по методическим указаниям, РПД, ГОСТам или учебникам ГУАП.",
                "is_out_of_scope": True,
                "sources": [],
            }

        # 1. Guardrail Scope Check
        scope_res = await self.check_scope(cleaned_q)
        if not scope_res.get("is_in_scope", True):
            rejection_text = (
                "Извините, я — специализированный агент-методист кафедры физики ГУАП. "
                "Моя область знаний строго ограничена учебными программами (РПД), методическими указаниями, "
                "учебниками Коваленко, задачниками, ГОСТами оформления студенческих работ и регламентами ГУАП. "
                "Я не уполномочен отвечать на вопросы вне этой предметной области."
            )
            await self.db.log_methodist_qa(
                caller_type=caller,
                caller_id=caller_id,
                question=cleaned_q,
                answer=rejection_text,
                sources_json=json.dumps([]),
                is_out_of_scope=True,
            )
            return {
                "answer": rejection_text,
                "is_out_of_scope": True,
                "sources": [],
                "scope_reason": scope_res.get("reason", "Out of domain"),
            }

        # 2. Search TOC Index for target book chapters (Table of Contents FTS)
        toc_matches = await self.db.search_methodist_toc(cleaned_q, limit=3)
        ephemeral_slices: list[dict[str, Any]] = []

        if toc_matches:
            # Load ephemeral slice for top 1-2 most relevant chapters
            for match in toc_matches[:2]:
                f_path = match.get("file_path")
                p_start = match.get("page_start", 1)
                p_end = match.get("page_end", p_start)
                b_title = match.get("book_title", "")
                ch_title = match.get("chapter_title", "")

                if f_path:
                    # Read max 6 pages per chapter slice
                    slice_content = self.get_ephemeral_pages(f_path, p_start, p_end, max_pages=6)
                    ephemeral_slices.append({
                        "book_title": b_title,
                        "chapter_title": ch_title,
                        "page_start": p_start,
                        "page_end": p_end,
                        "content": slice_content,
                    })

        # 3. Retrieve standard rules and guidelines from FTS chunks
        retrieved_chunks = await self.db.search_methodist_fts(cleaned_q, limit=4)

        # 4. Assemble Ephemeral Context
        context_blocks = []
        sources = []

        # Add ephemeral textbook/manual slices
        for s in ephemeral_slices:
            header = f"### [УЧЕБНОЕ ПОСОБИЕ]: {s['book_title']} — {s['chapter_title']} (стр. {s['page_start']}-{s['page_end']})"
            context_blocks.append(f"{header}\n{s['content']}")
            sources.append({
                "source_type": "textbook_toc",
                "doc_title": s["book_title"],
                "section": s["chapter_title"],
                "pages": f"{s['page_start']}-{s['page_end']}",
            })

        # Add markdown & template chunks
        for ch in retrieved_chunks:
            sec = ch.get("section_title", "")
            d_title = ch.get("doc_title", "")
            txt = ch.get("chunk_text", "")
            context_blocks.append(f"### [РЕГЛАМЕНТ/ШАБЛОН]: {d_title} [{sec}]\n{txt}")
            sources.append({
                "source_type": "chunk",
                "doc_title": d_title,
                "section": sec,
            })

        combined_context = "\n\n".join(context_blocks)
        if not combined_context:
            combined_context = "В базе знаний кафедры нет точных текстуальных совпадений. Отвечай на основе утвержденных стандартов кафедры физики ГУАП."

        # 5. Generate Response via Gemini (Isolated turn, base prompt stays frozen)
        user_prompt = f"""КОНТЕКСТ МЕТОДИЧЕСКОЙ БАЗЫ ЗНАНИЙ (ТОЧЕЧНЫЕ ВЫГРУЗКИ СТРАНИЦ И РЕГЛАМЕНТЫ):
{combined_context}

ВОПРОС:
{cleaned_q}

Сформируй исчерпывающий, академически строгий и точный ответ. Обязательно сошлись на конкретные разделы учебника Коваленко, сборника задач или методички к лабораторной работе (с указанием страниц/формул/номеров задач), если они представлены в контексте."""

        import asyncio
        client = self.get_client()
        answer_text = ""
        candidate_models = [self.model, "gemini-3.8-flash", "gemini-3.5-flash-lite", "gemini-3.5-flash", "gemini-flash-lite-latest"]
        # Deduplicate while preserving order
        unique_models = []
        for m in candidate_models:
            if m not in unique_models:
                unique_models.append(m)

        system_instruction = render_prompt("methodist")
        for attempt, mod in enumerate(unique_models):
            try:
                # Prefer LOW thinking level, or omit for older models
                cfg_kwargs: dict[str, Any] = {
                    "system_instruction": system_instruction,
                    "temperature": 0.2,
                }
                if "3.8" in mod or "3.7" in mod or "2.5" in mod:
                    cfg_kwargs["thinking_config"] = types.ThinkingConfig(thinking_level="LOW")

                resp = await asyncio.wait_for(
                    client.aio.models.generate_content(
                        model=mod,
                        contents=user_prompt,
                        config=types.GenerateContentConfig(**cfg_kwargs),
                    ),
                    timeout=25.0,
                )
                answer_text = (resp.text or "").strip()
                if answer_text:
                    logger.info("Successfully generated methodist answer via %s (attempt %d)", mod, attempt + 1)
                    break
            except Exception as e:
                logger.warning("Attempt %d with model %s failed: %s", attempt + 1, mod, e)
                if attempt < len(unique_models) - 1:
                    await asyncio.sleep(1.0)
                else:
                    answer_text = f"⚠️ Временная ошибка сервиса генерации методического ответа ({e}). Пожалуйста, повторите запрос через минуту."

        # 6. Context Wipe: Discard ephemeral slices immediately
        ephemeral_slices.clear()
        combined_context = ""

        # 7. Log interaction into persistent QA memory
        await self.db.log_methodist_qa(
            caller_type=caller,
            caller_id=caller_id,
            question=cleaned_q,
            answer=answer_text,
            sources_json=json.dumps(sources, ensure_ascii=False),
            is_out_of_scope=False,
        )

        return {
            "answer": answer_text,
            "is_out_of_scope": False,
            "sources": sources,
            "toc_matches_used": len(toc_matches),
            "chunks_used": len(retrieved_chunks),
        }

    # --- 4. Teaching & Catalog Inspector ---

    async def list_books(self) -> list[dict[str, Any]]:
        """List all indexed books and manuals with TOC statistics."""
        return await self.db.list_methodist_books_with_toc()

    async def get_book_toc(self, doc_id: int) -> list[dict[str, Any]]:
        """Get structured Table of Contents for a specific document."""
        return await self.db.get_methodist_toc_for_doc(doc_id)

    async def read_book_pages(self, doc_id: int, page_start: int, page_end: int) -> str:
        """On-demand inspection of specific pages in an indexed book."""
        books = await self.db.list_methodist_documents()
        doc = next((b for b in books if b["id"] == doc_id), None)
        if not doc:
            return f"Документ с ID {doc_id} не найден."

        f_path = doc.get("file_path")
        return self.get_ephemeral_pages(f_path, page_start, page_end, max_pages=10)

    async def teach_guideline(self, title: str, content: str, category: str = "dept_physics") -> dict[str, Any]:
        """Save a new guideline/rule into agent's directory, index it, and confirm learning."""
        target_dir = self.raw_dir / category
        target_dir.mkdir(parents=True, exist_ok=True)

        clean_slug = re.sub(r"[^\w\-]+", "_", title.lower()).strip("_")
        file_path = target_dir / f"{clean_slug}.md"

        header = f"# {title}\n\n*Добавлено преподавателем: {datetime.now().strftime('%Y-%m-%d %H:%M')}*\n\n"
        file_path.write_text(header + content.strip() + "\n", encoding="utf-8")

        ingest_res = await self.ingest_markdown_file(file_path, category=category)
        return {
            "status": "success",
            "file_path": str(file_path),
            "ingest_result": ingest_res,
        }

    async def get_stats(self) -> dict[str, Any]:
        """Return summary of indexed knowledge base including TOC entries."""
        docs = await self.db.list_methodist_documents()
        books_with_toc = await self.db.list_methodist_books_with_toc()

        by_cat = {}
        for d in docs:
            c = d.get("category", "other")
            by_cat[c] = by_cat.get(c, 0) + 1

        total_toc_entries = sum(b.get("toc_entries_count", 0) for b in books_with_toc)

        return {
            "total_documents": len(docs),
            "total_toc_entries": total_toc_entries,
            "by_category": by_cat,
            "books": [
                {
                    "id": b["doc_id"],
                    "title": b["title"],
                    "category": b["category"],
                    "toc_count": b["toc_entries_count"],
                    "pages": f"{b['min_page'] or 1}-{b['max_page'] or '?'}",
                }
                for b in books_with_toc
                if b.get("toc_entries_count", 0) > 0
            ],
        }
